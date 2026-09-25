"""Fills an HR-uploaded .docx offer-letter template and converts it to PDF.

Uses `docxtpl` (Jinja2-in-Word) to substitute `{{placeholder}}` tokens typed
directly into the document — it handles Word's habit of splitting one piece
of visible text across several XML runs, which naive find/replace on the
raw XML does not. The filled .docx is then converted to PDF by driving the
real MS Word installed on this machine via COM automation (`docx2pdf`).

This only works on Windows with MS Word installed — there is no portable
equivalent without LibreOffice's headless `--convert-to pdf` (see the note
in requirements.txt). If this backend ever needs to run on Linux/Mac or a
machine without Word, swap this module for a LibreOffice-based one; nothing
outside this file and services.py needs to change.
"""
import logging
import tempfile
from pathlib import Path

from docxtpl import DocxTemplate

logger = logging.getLogger(__name__)


class DocxRenderError(ValueError):
    """Raised when the uploaded template can't be filled or converted —
    e.g. it isn't a valid .docx, or Word/COM isn't available on this
    machine. Subclasses ValueError so it's caught by the same handling as
    other offer-letter validation errors (see serializers.py, views.py)
    rather than surfacing as a raw 500."""


def _convert_via_word(docx_path: str, pdf_path: str) -> None:
    """`docx2pdf.convert()` drives Word over COM but never calls
    `CoInitialize()` itself — fine for a one-off script's main thread, but
    Django's dev/prod server handles each request on a pooled worker thread
    that has never touched COM, which raises `CoInitialize has not been
    called` (confirmed live: this broke the very first request that landed
    on a fresh thread). Every COM-using thread needs its own init/uninit."""
    import pythoncom
    from docx2pdf import convert

    pythoncom.CoInitialize()
    try:
        convert(docx_path, pdf_path)
    finally:
        pythoncom.CoUninitialize()


def _append_signature_block(docx_path: Path, context: dict) -> None:
    """Mirrors `_render_text_template_pdf`'s post-signature block for the
    reportlab path (offer_letter.py) — an uploaded .docx template has no
    `{{signature_name}}`/`{{signed_date}}` tokens of its own (HR never typed
    them in), so docxtpl's substitution silently has nothing to fill. This
    appends the same "Signed and Accepted" section directly, so a signed
    offer looks signed regardless of which renderer produced it."""
    from docx import Document

    document = Document(str(docx_path))
    document.add_paragraph()
    document.add_heading('Signed and Accepted', level=2)
    document.add_paragraph(f"Signature: {context.get('signature_name', '')}")
    document.add_paragraph(f"Date: {context.get('signed_date', '')}")
    document.save(str(docx_path))


def render_docx_template_to_pdf(source_path: str, context: dict, *, signed: bool = False) -> bytes:
    with tempfile.TemporaryDirectory(prefix='offer_letter_') as tmp_dir:
        tmp_dir_path = Path(tmp_dir)
        try:
            tpl = DocxTemplate(source_path)
            tpl.render(context)
        except Exception as exc:  # docxtpl/Jinja2 raise several distinct types
            raise DocxRenderError(f'Could not fill the Word template: {exc}') from exc

        rendered_docx = tmp_dir_path / 'rendered.docx'
        tpl.save(str(rendered_docx))

        if signed:
            _append_signature_block(rendered_docx, context)

        rendered_pdf = tmp_dir_path / 'rendered.pdf'
        try:
            _convert_via_word(str(rendered_docx), str(rendered_pdf))
        except ImportError as exc:
            logger.exception('docx2pdf/pywin32 not importable')
            raise DocxRenderError(
                'Could not convert the Word template to PDF: the server is missing the '
                f'"{exc.name}" Python package. Run `pip install -r requirements.txt` with the '
                'same Python that runs the backend, then restart it.'
            ) from exc
        except Exception as exc:
            logger.exception('docx2pdf conversion failed')
            raise DocxRenderError(
                'Could not convert the Word template to PDF. This requires Microsoft Word '
                'to be installed on the server — check it is available.'
            ) from exc

        if not rendered_pdf.exists():
            raise DocxRenderError('Word did not produce a PDF file — conversion failed silently.')

        return rendered_pdf.read_bytes()
