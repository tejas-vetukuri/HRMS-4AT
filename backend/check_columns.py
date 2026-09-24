import openpyxl

wb = openpyxl.load_workbook("roster_temp.xlsx")
ws = wb.active

print(f"Total rows: {ws.max_row}")
print("\n=== Checking first 10 rows ===\n")

for row_num in range(1, 11):
    row = ws[row_num]
    values = [cell.value for cell in row[:5]]
    print(f"Row {row_num}: {values}")

